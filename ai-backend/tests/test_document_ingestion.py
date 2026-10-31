import base64
import io
import os
import zipfile

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from kb.documents import CHUNK_SIZE, chunk_text, extract_text


def make_pdf(text="Apple harvest begins in September."):
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'),
                             NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
    stream = DecodedStreamObject()
    stream.set_data(f'BT /F1 12 Tf 20 250 Td ({text}) Tj ET'.encode('ascii'))
    page[NameObject('/Contents')] = writer._add_object(stream)
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def make_docx():
    doc = Document()
    doc.add_paragraph('First paragraph')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = 'blue'
    table.cell(0, 1).text = 'whale'
    doc.add_paragraph('Last paragraph')
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def upload_bytes(kb, name, raw):
    return kb.client.post('/api/kb/upload', json={'filename': name, 'content_b64': base64.b64encode(raw).decode()})


def test_docx_extracts_paragraphs_and_table_in_order(kb):
    response = upload_bytes(kb, 'a.docx', make_docx())
    assert response.status_code == 200
    text = kb.client.get(f"/api/kb/documents/{response.json()['doc_id']}/content").json()['content']
    assert text == 'First paragraph\n\nblue | whale\n\nLast paragraph'
    assert 'blue | whale' in kb.vectors.insert_chunks.call_args.args[1][0]


def test_pdf_text_upload(kb):
    response = upload_bytes(kb, 'a.pdf', make_pdf())
    assert response.status_code == 200
    assert 'Apple harvest' in kb.client.get(f"/api/kb/documents/{response.json()['doc_id']}/content").json()['content']


@pytest.mark.parametrize('name,raw', [('a.pdf', b'bad pdf'), ('a.docx', b'bad zip'), ('a.doc', b'hi'), ('noextension', b'hi')])
def test_invalid_files_rejected(kb, name, raw):
    assert upload_bytes(kb, name, raw).status_code == 400
    kb.vectors.insert_chunks.assert_not_called()


def test_docx_zip_expansion_limit(kb):
    out = io.BytesIO()
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as z:
        z.writestr('huge.xml', b'x' * (21 * 1024 * 1024))
    assert upload_bytes(kb, 'a.docx', out.getvalue()).status_code == 400
    kb.vectors.insert_chunks.assert_not_called()


def test_scanned_pdf_rejected_without_ocr(kb):
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    out = io.BytesIO()
    writer.write(out)
    assert upload_bytes(kb, 'scan.pdf', out.getvalue()).status_code == 400
    kb.vectors.insert_chunks.assert_not_called()


def test_pdf_with_text_page_and_textless_page_is_rejected(kb):
    writer = PdfWriter()
    writer.append(io.BytesIO(make_pdf()))
    writer.add_blank_page(width=300, height=300)
    out = io.BytesIO()
    writer.write(out)
    response = upload_bytes(kb, 'mixed.pdf', out.getvalue())
    assert response.status_code == 400
    assert '无法提取文字' in response.json()['detail']
    kb.vectors.insert_chunks.assert_not_called()
    assert kb.client.get('/api/kb/documents').json() == []


def test_paragraph_boundaries_and_long_text():
    text = 'A' * 250 + '\n\n' + 'B' * 250 + '\n\n' + 'C' * 1000
    chunks = chunk_text(text, 'txt')
    assert all(0 < len(c) <= CHUNK_SIZE for c in chunks)
    assert any('A' * 250 in c for c in chunks)
    assert any('B' * 250 in c for c in chunks)
    assert sum(c.count('C') for c in chunks) >= 1000


def test_markdown_fenced_block_preserves_blank_lines():
    text = 'intro\n\n```python\na = 1\n\nb = 2\n```\n\noutro'
    assert '```python\na = 1\n\nb = 2\n```' in chunk_text(text, 'md')[0]


@pytest.mark.skipif(os.environ.get('RUN_RAG_MODEL_TESTS') != '1', reason='requires local real embedding model')
def test_real_model_synthetic_recall_ci(tmp_path):
    from eval.retrieval_benchmark import evaluate
    doc = make_docx()
    documents = [('notes.txt', b'The observatory telescope is on the mountain.'),
                 ('guide.md', b'# Kitchen\n\nThe orange soup recipe uses carrots.'),
                 ('manual.docx', doc), ('report.pdf', make_pdf())]
    for name, raw in documents:
        (tmp_path / name).write_bytes(raw)
    manifest = {'documents': [{'id': n, 'path': str(tmp_path / n)} for n, _ in documents],
                'queries': [
                    {'id': 'txt', 'document': 'notes.txt', 'query': 'Where is the telescope?', 'evidence': 'telescope is on the mountain'},
                    {'id': 'md', 'document': 'guide.md', 'query': 'What ingredient is in the soup?', 'evidence': 'soup recipe uses carrots'},
                    {'id': 'docx', 'document': 'manual.docx', 'query': 'What animal is blue?', 'evidence': 'blue | whale'},
                    {'id': 'pdf', 'document': 'report.pdf', 'query': 'When is the apple harvest?', 'evidence': 'Apple harvest begins in September'},
                ]}
    results = evaluate(manifest, tmp_path / 'milvus.db')
    recall = sum(r['selected_at_8'] for r in results) / len(results)
    print(f'\nSynthetic selected evidence recall@8: {recall:.0%} ({sum(r["selected_at_8"] for r in results)}/{len(results)})')
    assert recall >= 0.75
